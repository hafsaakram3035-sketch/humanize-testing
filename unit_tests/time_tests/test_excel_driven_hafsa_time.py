"""Excel + CSV driven tests for the time module (Task 3, Hafsa).

The test cases are written in assignment_docs/test_cases_hafsa.xlsx.
Each row names a module, a humanize function and a CSV file in test_data/.
This file only runs the rows whose Module column is "time".
"""
import csv
from pathlib import Path

import humanize
import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
EXCEL_FILE = ROOT / "assignment_docs" / "test_cases_hafsa.xlsx"
DATA_DIR = ROOT / "test_data"
MODULE = "time"

pytestmark = [pytest.mark.datadriven, pytest.mark.excel]


def load_cases():
    """Build the list of test cases from the Excel sheet and the CSV files."""
    workbook = load_workbook(EXCEL_FILE, data_only=True)
    sheet = workbook["TestCases"]
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()

    header = [str(name).strip() for name in rows[0]]
    cases = []
    for row in rows[1:]:
        if not row[0]:
            continue  # skip empty lines
        info = dict(zip(header, row))
        if str(info["Module"]).strip().lower() != MODULE:
            continue  # this row belongs to the other module
        test_id = str(info["Test ID"]).strip()
        function = getattr(humanize, str(info["Function"]).strip().lower())
        csv_path = DATA_DIR / str(info["Data File"]).strip()
        with open(csv_path, newline="", encoding="utf-8") as f:
            for data in csv.DictReader(f):
                cases.append(pytest.param(
                    function, int(data["input"]), data["expected"],
                    id=f"{test_id}-{data['input']}",
                ))
    return cases


@pytest.mark.parametrize("function, value, expected", load_cases())
def test_function_from_excel_and_csv(function, value, expected):
    assert function(value) == expected